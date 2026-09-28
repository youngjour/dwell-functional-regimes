"""Step 19 - validate HSMM regime meaning against Seoul's official 121 major places
and check emergent activity cores vs real commercial/tourism areas. No model rerun.

Places shapefile CRS = WGS84 (EPSG:4326, .prj explicit) -> reproject to EPSG:5179.
dbf is UTF-8 (.cpg correct); AREA_CD (ASCII) joins to the xlsx for clean CATEGORY/
AREA_NM. Cells (centroids, 5179) -> place polygons (within; polygons may overlap so a
cell can map to multiple places).

Category / place names are kept in Korean as in the source xlsx: 발달상권 (commercial
district), 관광특구 (special tourist zone), 인구밀집지역 (population-concentrated area),
공원 (park), 고궁·문화유산 (palace / cultural heritage).

Outputs: data/processed/place_cell_map.parquet, results/stats/place_validation_stats.json.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import polars as pl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config as C  # noqa: E402

PLACES_SHP = (C.PROCESSED_DIR / "places" / "서울시 주요 121장소 영역" /
              "서울시 주요 121장소 영역.shp")
XLSX = C.PLAN_RAW_DIR / "서울시 주요 121장소 목록.xlsx"
ACT = ["activity_A", "activity_B"]
REGIMES = ["residential", "activity_A", "activity_B", "dense_mixed",
           "mid_density", "low_density", "structural_missing"]
CAT_ORDER = ["발달상권", "관광특구", "인구밀집지역", "공원", "고궁·문화유산"]


def load_places():
    import geopandas as gpd
    import openpyxl
    g = gpd.read_file(PLACES_SHP)[["AREA_CD", "geometry"]]      # AREA_CD ASCII-safe
    g = g.set_crs(4326, allow_override=True).to_crs(5179)
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2) if r[2].value]
    meta = pl.DataFrame({"AREA_CD": [r[2] for r in rows],
                         "CATEGORY": [r[0] for r in rows],
                         "AREA_NM": [r[3] for r in rows]})
    return g, meta


def cells_gdf(reg):
    import geopandas as gpd
    from shapely.geometry import Point
    pdf = reg.to_pandas()
    return gpd.GeoDataFrame(pdf, geometry=[Point(x, y) for x, y in
                            zip(reg["utm_x"], reg["utm_y"])], crs=5179)


def main():
    import geopandas as gpd
    reg = pl.read_parquet(C.PROCESSED_DIR / "cell_regime_summary.parquet").with_columns(
        (pl.col("share_activity_A") + pl.col("share_activity_B")).alias("activity_share"))
    plan = pl.read_parquet(C.PROCESSED_DIR / "cell_plan_map.parquet").select(
        ["cell_id", "jungsim", "jungsim_tier"])
    reg = reg.join(plan, on="cell_id")
    cg = cells_gdf(reg)

    places, meta = load_places()
    # CRS sanity vs grid
    gb = gpd.read_file(C.GRID_SHP).total_bounds
    pb = places.total_bounds
    overlap_ok = (pb[0] > gb[0] - 5000 and pb[2] < gb[2] + 5000 and
                  pb[1] > gb[1] - 5000 and pb[3] < gb[3] + 5000)
    print(f"CRS overlay ok={overlap_ok} place_bounds={[round(v) for v in pb]}", flush=True)

    # cell -> place (within; multi-membership allowed)
    j = gpd.sjoin(cg[["cell_id", "geometry"]], places, predicate="within", how="inner")
    pcm = pl.DataFrame({"cell_id": j["cell_id"].tolist(),
                        "AREA_CD": j["AREA_CD"].tolist()}).join(meta, on="AREA_CD")
    pcm = pcm.join(reg.select(["cell_id", "dominant_functional", "activity_share",
                               "share_residential", "jungsim"]), on="cell_id")
    pcm.write_parquet(C.PROCESSED_DIR / "place_cell_map.parquet")
    stats = {"crs": "EPSG:4326->5179", "crs_overlay_ok": bool(overlap_ok),
             "n_places": meta.height, "category_counts": meta["CATEGORY"].value_counts().to_dict(as_series=False)}

    # ---- 2. CATEGORY x regime composition ----
    cat_rows = []
    M = np.zeros((len(CAT_ORDER), len(REGIMES)))
    for ci, cat in enumerate(CAT_ORDER):
        sub = pcm.filter(pl.col("CATEGORY") == cat)
        ncell = sub["cell_id"].n_unique()
        if sub.height == 0:
            continue
        domc = sub.group_by("dominant_functional").len()
        dm = {r["dominant_functional"]: r["len"] for r in domc.iter_rows(named=True)}
        tot = sum(dm.values())
        for ri, rg in enumerate(REGIMES):
            M[ci, ri] = dm.get(rg, 0) / tot
        cat_rows.append({
            "category": cat, "n_places": meta.filter(pl.col("CATEGORY") == cat).height,
            "n_cells": int(ncell),
            "mean_activity_share": round(float(sub["activity_share"].mean()), 3),
            "mean_residential_share": round(float(sub["share_residential"].mean()), 3),
            "pct_dom_activity": round(100 * sub["dominant_functional"].is_in(ACT).mean(), 1),
            "pct_dom_residential": round(100 * (sub["dominant_functional"] == "residential").mean(), 1)})
    stats["category_regime"] = cat_rows
    stats["category_dom_matrix"] = {"rows": CAT_ORDER, "cols": REGIMES, "M": M.tolist()}

    # ---- 3. emergent validation ----
    emergent = reg.filter(pl.col("dominant_functional").is_in(ACT) & pl.col("jungsim").is_null())
    em_cells = set(emergent["cell_id"].to_list())
    comm = pcm.filter(pl.col("CATEGORY").is_in(["발달상권", "관광특구"]))
    comm_cells = set(comm["cell_id"].to_list())
    em_in_comm = len(em_cells & comm_cells)
    stats["emergent_validation"] = {
        "n_emergent": len(em_cells),
        "n_in_commercial_tourism": em_in_comm,
        "pct_in_commercial_tourism": round(100 * em_in_comm / len(em_cells), 1),
        "n_in_any_place": len(em_cells & set(pcm["cell_id"].to_list())),
        "pct_in_any_place": round(100 * len(em_cells & set(pcm["cell_id"].to_list())) / len(em_cells), 1)}

    # per-cluster: which place contains the cluster centroid
    clusters = pl.read_parquet(C.PROCESSED_DIR / "emergent_activity_clusters.parquet")
    from shapely.geometry import Point
    cl_rows = []
    for r in clusters.iter_rows(named=True):
        pt = Point(r["cx"], r["cy"])
        hit = places[places.contains(pt)]
        nm = cat = None
        if len(hit):
            cd = hit.iloc[0]["AREA_CD"]
            mr = meta.filter(pl.col("AREA_CD") == cd)
            if mr.height:
                nm = mr["AREA_NM"][0]; cat = mr["CATEGORY"][0]
        cl_rows.append({"cluster_id": r.get("cluster_id"),
                        "n_cells": r["n_cells"], "dominant_admin_dong": r["dominant_admin_dong"],
                        "mean_activity_share": r["mean_activity_share"],
                        "place_name": nm, "place_category": cat})
    stats["emergent_clusters_placed"] = cl_rows
    n_cl_placed = sum(1 for c in cl_rows if c["place_name"])
    stats["emergent_clusters_in_place_pct"] = round(100 * n_cl_placed / len(cl_rows), 1)

    # ---- reverse: commercial districts not in any designated center (gap candidates) ----
    sangkwon = pcm.filter(pl.col("CATEGORY") == "발달상권")
    per_sk = sangkwon.group_by("AREA_NM").agg(
        pl.len().alias("n_cells"),
        pl.col("jungsim").is_not_null().mean().alias("frac_in_jungsim"),
        pl.col("activity_share").mean().alias("mean_act"),
        pl.col("dominant_functional").is_in(ACT).mean().alias("frac_act_dom"))
    gaps = per_sk.filter(pl.col("frac_in_jungsim") < 0.1).sort("n_cells", descending=True)
    stats["sangkwon_planning_gaps"] = {
        "n_sangkwon_total": sangkwon["AREA_NM"].n_unique(),
        "n_gap_(no_jungsim)": gaps.height,
        "gap_list": gaps.select(["AREA_NM", "n_cells", "mean_act", "frac_act_dom"]).to_dicts()}

    # ---- 4. commercial/tourism places vs designated-center overlap ----
    ct = pcm.filter(pl.col("CATEGORY").is_in(["발달상권", "관광특구"]))
    stats["commercial_in_jungsim_pct"] = round(
        100 * ct.filter(pl.col("jungsim").is_not_null()).height / ct.height, 1)

    (C.STATS_DIR / "place_validation_stats.json").write_text(
        json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
    print("DONE")


if __name__ == "__main__":
    main()
